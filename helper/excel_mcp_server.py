# helper/excel_mcp_server.py
import os
import io
import contextlib
import traceback
import openpyxl
from openpyxl.cell import Cell
from openpyxl.utils import get_column_letter
import pandas as pd
from fastmcp import FastMCP

mcp = FastMCP("ExcelUniversalQueryServer")


def _resolve_path(file_path: str) -> str:
    """Tìm đường dẫn file chính xác, hỗ trợ các thư mục đầu vào khác nhau."""
    if os.path.exists(file_path):
        return file_path
    for folder in ["excel_input", "xlsx_files", "/input/xlsx_files"]:
        alt_path = os.path.join(folder, os.path.basename(file_path))
        if os.path.exists(alt_path):
            return alt_path
    raise FileNotFoundError(f"Không tìm thấy tệp Excel: {file_path}")


def _unmerge_and_fill(sheet):
    """Tách toàn bộ ô gộp (merged cells) trên RAM và điền giá trị ô gốc vào các ô con."""
    for m_range in list(sheet.merged_cells.ranges):
        top_left_val = sheet.cell(m_range.min_row, m_range.min_col).value
        sheet.unmerge_cells(str(m_range))
        for r in range(m_range.min_row, m_range.max_row + 1):
            for c in range(m_range.min_col, m_range.max_col + 1):
                cell = sheet.cell(r, c)
                if isinstance(cell, Cell):
                    cell.value = top_left_val


@mcp.tool()
def get_workbook_overview(file_path: str, max_preview_rows: int = 40) -> str:
    """
    Đọc cấu trúc tổng quan của toàn bộ tệp Excel (.xlsx):
    Liệt kê tất cả các sheet, kích thước dòng/cột, tự động làm phẳng ô gộp (merged cells)
    và trả về bảng dữ liệu xem trước kèm tọa độ Dòng (R1, R2...) và Cột (C1(A), C2(B)...).
    """
    try:
        path = _resolve_path(file_path)
        wb = openpyxl.load_workbook(path, data_only=True)
        out = [
            f"=== WORKBOOK: {os.path.basename(path)} | Danh sách Sheets: {wb.sheetnames} ==="]

        for name in wb.sheetnames:
            sheet = wb[name]
            merged_count = len(sheet.merged_cells.ranges)
            _unmerge_and_fill(sheet)

            max_r = min(max_preview_rows, sheet.max_row)
            max_c = min(20, sheet.max_column)
            out.append(
                f"\n--- SHEET: '{name}' (Kích thước: {sheet.max_row} dòng x {sheet.max_column} cột | Ô gộp đã tách: {merged_count}) ---"
            )

            headers = [
                "Row"] + [f"C{c}({get_column_letter(c)})" for c in range(1, max_c + 1)]
            rows_data = []
            for r in range(1, max_r + 1):
                row_vals = [sheet.cell(
                    r, c).value for c in range(1, max_c + 1)]
                if any(v is not None for v in row_vals):
                    clean_vals = [
                        str(v).replace("\n", " ").strip()[
                            :65] if v is not None else ""
                        for v in row_vals
                    ]
                    rows_data.append([f"R{r}"] + clean_vals)

            if rows_data:
                df = pd.DataFrame(rows_data, columns=headers)
                out.append(df.to_markdown(index=False))
                if sheet.max_row > max_preview_rows:
                    out.append(
                        f"... (Còn {sheet.max_row - max_preview_rows} dòng bên dưới. Dùng read_sheet_grid hoặc run_python_analysis để xử lý toàn bộ)"
                    )
            else:
                out.append("(Sheet trống)")

        wb.close()
        return "\n".join(out)
    except Exception as e:
        return f"Lỗi khi đọc tổng quan workbook: {str(e)}"


@mcp.tool()
def read_sheet_grid(
    file_path: str,
    sheet_name: str = "",
    start_row: int = 1,
    end_row: int = 120,
    start_col: int = 1,
    end_col: int = 25,
    fill_merged: bool = True,
) -> str:
    """
    Đọc chi tiết một vùng dòng/cột của sheet dưới dạng bảng Markdown có tọa độ R (dòng) và C (cột).
    """
    try:
        path = _resolve_path(file_path)
        wb = openpyxl.load_workbook(path, data_only=True)
        if not sheet_name or sheet_name not in wb.sheetnames:
            sheet_name = wb.sheetnames[0]
        sheet = wb[sheet_name]

        if fill_merged:
            _unmerge_and_fill(sheet)

        max_r = min(end_row, sheet.max_row)
        max_c = min(end_col, sheet.max_column)
        headers = [
            "Row"] + [f"C{c}({get_column_letter(c)})" for c in range(start_col, max_c + 1)]
        rows_data = []

        for r in range(start_row, max_r + 1):
            row_vals = [sheet.cell(r, c).value for c in range(
                start_col, max_c + 1)]
            if any(v is not None for v in row_vals):
                formatted_vals = [
                    str(v).replace("\n", " ").strip() if v is not None else ""
                    for v in row_vals
                ]
                rows_data.append([f"R{r}"] + formatted_vals)

        wb.close()
        if not rows_data:
            return f"Sheet '{sheet_name}' không có dữ liệu trong dòng {start_row}-{max_r}."

        df = pd.DataFrame(rows_data, columns=headers)
        return f"Sheet: '{sheet_name}' (Dòng {start_row} -> {max_r}):\n" + df.to_markdown(index=False)
    except Exception as e:
        return f"Lỗi khi đọc sheet: {str(e)}"


@mcp.tool()
def search_in_workbook(file_path: str, keyword: str) -> str:
    """Tìm kiếm nhanh vị trí (Sheet, Dòng, Cột) của một từ khóa trên toàn bộ các sheet."""
    try:
        path = _resolve_path(file_path)
        wb = openpyxl.load_workbook(path, data_only=True)
        matches = []
        kw_lower = keyword.lower()

        for name in wb.sheetnames:
            sheet = wb[name]
            for r in range(1, sheet.max_row + 1):
                for c in range(1, sheet.max_column + 1):
                    val = sheet.cell(r, c).value
                    if val is not None and kw_lower in str(val).lower():
                        col_letter = get_column_letter(c)
                        matches.append(
                            f"[{name}] R{r}C{c} ({col_letter}{r}): {repr(val)}")
                        if len(matches) >= 50:
                            wb.close()
                            return "Tìm thấy >50 kết quả (hiển thị 50 kết quả đầu):\n" + "\n".join(matches)
        wb.close()
        return "\n".join(matches) if matches else f"Không tìm thấy '{keyword}'."
    except Exception as e:
        return f"Lỗi khi tìm kiếm: {str(e)}"


@mcp.tool()
def run_python_analysis(file_path: str, python_code: str) -> str:
    """
    Thực thi code Python để phân tích và tính toán số liệu chính xác từ file Excel.
    Biến có sẵn:
    - `file_path`: đường dẫn thực tế tới file Excel.
    - `pd`: thư viện pandas.
    - `openpyxl`: thư viện openpyxl.
    - `load_unmerged_df(sheet_name=None)`: hàm trả về pandas DataFrame của sheet đã tự động tách và điền đầy mọi ô gộp (merged cells), index dòng và cột bắt đầu từ 0 (tương ứng R1->index 0, C1->col 0).
    BẮT BUỘC dùng `print()` để xuất kết quả tính toán.
    """
    try:
        resolved_path = _resolve_path(file_path)
        stdout_capture = io.StringIO()

        def load_unmerged_df(sheet_name=None):
            wb = openpyxl.load_workbook(resolved_path, data_only=True)
            s_name = sheet_name if (
                sheet_name and sheet_name in wb.sheetnames) else wb.sheetnames[0]
            sheet = wb[s_name]
            _unmerge_and_fill(sheet)
            data = list(sheet.values)
            wb.close()
            return pd.DataFrame(data)

        exec_globals = {
            "__builtins__": __builtins__,
            "file_path": resolved_path,
            "pd": pd,
            "openpyxl": openpyxl,
            "load_unmerged_df": load_unmerged_df,
        }

        with contextlib.redirect_stdout(stdout_capture):
            exec(python_code, exec_globals)

        output = stdout_capture.getvalue().strip()
        return output if output else "Code chạy thành công nhưng không có lệnh print() xuất kết quả."
    except Exception:
        return f"Lỗi khi thực thi Python code:\n{traceback.format_exc()}"


if __name__ == "__main__":
    mcp.run(transport="stdio")
